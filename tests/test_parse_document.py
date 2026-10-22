import pytest
from pptx import Presentation
from reportlab.pdfgen import canvas

from sift.parsing import parse_document


def test_parse_document_picks_parser_by_extension(tmp_path):
    pdf_path = tmp_path / "a.pdf"
    c = canvas.Canvas(str(pdf_path))
    c.drawString(100, 750, "Hello")
    c.showPage()
    c.save()

    pptx_path = tmp_path / "b.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Hello"
    prs.save(pptx_path)

    assert parse_document(pdf_path)[0].text == "Hello"
    assert parse_document(pptx_path)[0].text == "Hello"

def test_parse_document_rejects_unknown_extension(tmp_path):
    path = tmp_path / "notes.xyz"
    path.write_text("x")

    with pytest.raises(ValueError):
        parse_document(path)