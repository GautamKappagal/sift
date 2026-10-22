from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from sift.parsing.pptx_parser import parse_pptx


def test_parse_pptx_reads_title_and_bullets(tmp_path):
    path = tmp_path / "os_lec03.pptx"

    prs = Presentation()
    slide_layout = prs.slide_layouts[1]  # Title and Content layout
    slide = prs.slides.add_slide(slide_layout)

    slide.shapes.title.text = "Deadlock"

    body = slide.placeholders[1].text_frame
    body.text = "Mutual exclusion"
    body.add_paragraph().text = "Hold and wait"

    prs.save(path)

    blocks = parse_pptx(path)

    assert [b.text for b in blocks] == ["Deadlock", "Mutual exclusion", "Hold and wait"]
    assert blocks[0].block_type == "title"
    assert blocks[1].id == "os_lec03-s1-b2"

def test_parse_pptx_links_nested_bullets_to_parent(tmp_path):
    path = tmp_path / "os_lec03.pptx"

    prs = Presentation()
    slide_layout = prs.slide_layouts[1]  # Title and Content layout
    slide = prs.slides.add_slide(slide_layout)
    body = slide.placeholders[1].text_frame

    body.text = "Deadlock requires four conditions"
    p = body.add_paragraph()
    p.text = "Mutual exclusion"
    p.level = 1

    p = body.add_paragraph()
    p.text = "Hold and wait"
    p.level = 1

    body.add_paragraph().text = "Example: two trains on one track"

    prs.save(path)

    blocks = parse_pptx(path)

    assert blocks[0].parent_id is None
    assert blocks[1].parent_id == blocks[0].id
    assert blocks[2].parent_id == blocks[0].id
    assert blocks[3].parent_id is None

def test_parse_pptx_restarts_block_numbering_on_each_slide(tmp_path):
    path = tmp_path / "os_lec03.pptx"

    prs = Presentation()
    slide_layout = prs.slide_layouts[1]  # Title and Content layout

    # Slide 1
    slide1 = prs.slides.add_slide(slide_layout)
    slide1.shapes.title.text = "Deadlock"
    slide1.placeholders[1].text_frame.text = "Mutual exclusion"

    # Slide 2
    slide2 = prs.slides.add_slide(slide_layout)
    slide2.shapes.title.text = "Banker's algorithm"
    slide2.placeholders[1].text_frame.text = "Safe state"

    prs.save(path)

    blocks = parse_pptx(path)

    #blocks is one flat list: [s1-title, s1-bullet, s2-title, s2-bullet]
    assert blocks[0].id == "os_lec03-s1-b1"
    assert blocks[1].id == "os_lec03-s1-b2"
    assert blocks[2].id == "os_lec03-s2-b1"
    assert blocks[3].id == "os_lec03-s2-b2"

def test_parse_pptx_extracts_images_as_blocks(tmp_path):
    img_path = tmp_path / "diagram.png"
    Image.new("RGB", (10, 10)).save(img_path)

    path = tmp_path / "os_lec03.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout
    slide.shapes.add_picture(str(img_path), Inches(1), Inches(1))
    prs.save(path)

    out_dir = tmp_path / "images"
    blocks = parse_pptx(path, image_dir=out_dir)

    assert len(blocks) == 1
    assert blocks[0].id == "os_lec03-s1-b1"
    assert blocks[0].block_type == "image"
    assert blocks[0].text == ""
    assert Path(blocks[0].metadata["image_path"]).exists()