from pathlib import Path

from pptx import Presentation

from sift.models import SlideBlock


def parse_pptx(path: Path) -> list[SlideBlock]:
    """Parse a PPTX file and extract slide blocks.

    Args:
        path (Path): The path to the PPTX file.

    Returns:
        list[SlideBlock]: A list of slide blocks extracted from the PPTX file.
    """
    deck = path.stem.lower()
    blocks = []

    for slide_number, slide in enumerate(Presentation(path).slides, start=1):
        # Extract content
        for shape in slide.shapes:
            if shape.has_text_frame:
                if shape == slide.shapes.title:
                    block = SlideBlock(
                        id=f"{deck}-s{slide_number}-b{len(blocks) + 1}",
                        slide_number=slide_number,
                        text=shape.text_frame.text.strip(),
                        block_type="title"
                    )

                    blocks.append(block)

                else:
                    for paragraph in shape.text_frame.paragraphs:
                        if paragraph.text.strip():
                            block = SlideBlock(
                                id=f"{deck}-s{slide_number}-b{len(blocks) + 1}",
                                slide_number=slide_number,
                                text=paragraph.text.strip(),
                                block_type="content"
                            )

                            blocks.append(block)

    return blocks